
import shutil
import tempfile

from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation

from extractors.common import logger
from extractors.image_extractor import image_extract


def normalize_image(image_bytes: bytes) -> bytes | None:
    try:
        image = Image.open(
            BytesIO(image_bytes)
        ).convert("RGB")

        output = BytesIO()

        image.save(output, format="PNG")

        return output.getvalue()

    except Exception as e:
        logger.warning(
            f"Image normalization failed: {e}"
        )

        return None


def analyze_image(
    image_bytes: bytes,
    page_number: int,
) -> str:

    try:
        temp_image = (
            Path(tempfile.mkdtemp())
            / f"slide_{page_number}.png"
        )

        with open(temp_image, "wb") as f:
            f.write(image_bytes)

        text = image_extract(temp_image)

        shutil.rmtree(
            temp_image.parent,
            ignore_errors=True,
        )

        return (
            f"\n[IMAGE OCR - SLIDE "
            f"{page_number}]\n{text}\n"
        )

    except Exception as e:

        return (
            f"\n[IMAGE OCR FAILED "
            f"ON SLIDE {page_number}: "
            f"{e}]\n"
        )


def pptx_extract(file_path: Path) -> str:

    try:
        logger.info("Extracting PPTX text...")

        prs = Presentation(file_path)

        text_parts = []

        for slide_index, slide in enumerate(prs.slides):

            text_parts.append(
                f"\n--- SLIDE "
                f"{slide_index + 1} ---\n"
            )

            for shape in slide.shapes:

                try:
                    if hasattr(shape, "text"):

                        text = shape.text.strip()

                        if text:
                            text_parts.append(text)

                    if not hasattr(shape, "image"):
                        continue

                    image_bytes = shape.image.blob

                    png_bytes = normalize_image(
                        image_bytes
                    )

                    if png_bytes is None:
                        continue

                    image_block = analyze_image(
                        image_bytes=png_bytes,
                        page_number=slide_index + 1,
                    )

                    text_parts.append(image_block)

                except Exception as e:
                    logger.warning(
                        f"PPTX image failed: {e}"
                    )

        return "\n".join(text_parts)

    except Exception as e:
        raise RuntimeError(
            f"PPTX extraction failed: {e}"
        ) from e